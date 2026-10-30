import os
import io
import base64
import json
from datetime import datetime
from functools import wraps

import psycopg2
from psycopg2.extras import DictCursor
from flask import Flask, request, redirect, render_template_string, send_file, jsonify, session
from werkzeug.security import generate_password_hash, check_password_hash

# --- LIBRERÍAS DE EXCEL Y PDF ---
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet

# --- IA ---
from groq import Groq

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "clave-super-secreta-liceo-2024")

# 🔥 CONEXIÓN A SUPABASE (POSTGRESQL)
DATABASE_URL = os.environ.get("DATABASE_URL")
if not DATABASE_URL:
    raise Exception("Falta la variable DATABASE_URL en Render")

def get_db():
    return psycopg2.connect(DATABASE_URL)

# SEGURIDAD DE LA API KEY
GROQ_API_KEY = os.environ.get("GROQ_API_KEY")

# ==========================================
# FUNCIONES DE BASE DE DATOS Y SEGURIDAD
# ==========================================
def init_db():
    con = get_db()
    cur = con.cursor()
    cur.execute("""CREATE TABLE IF NOT EXISTS usuarios(
        id SERIAL PRIMARY KEY,
        username TEXT UNIQUE NOT NULL,
        password TEXT NOT NULL)""")
    cur.execute("""CREATE TABLE IF NOT EXISTS docs(
        id SERIAL PRIMARY KEY,
        user_id INTEGER NOT NULL,
        tipo TEXT, tipo_doc TEXT, fecha TEXT, folio TEXT, rut TEXT, nombre TEXT,
        neto REAL, exento REAL, iva REAL, total REAL, pago TEXT, nota TEXT)""")
    con.commit(); cur.close(); con.close()

init_db()

def login_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if 'user_id' not in session:
            return redirect('/login')
        return f(*args, **kwargs)
    return decorated_function

# ==========================================
# DISEÑO DE LA APP (CON MODO OSCURO)
# ==========================================
HTML_BASE = """
<!doctype html><html><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1,maximum-scale=1">
<title>Mi Libro CV</title>
<style>
 :root { --bg: #e8f5e9; --text: #1b5e20; --card: #fff; --border: #c8e6c9; --primary: #2e7d32; --btn: #43a047; --input-bg: #f1f8e9; --table-th: #2e7d32; }
 body.dark-mode { --bg: #121212; --text: #e0e0e0; --card: #1e1e1e; --border: #333; --primary: #1b5e20; --btn: #2e7d32; --input-bg: #2c2c2c; --table-th: #1b5e20; }

 body{font-family:system-ui,sans-serif;margin:0;padding:0;background:var(--bg);color:var(--text);padding-bottom:70px; transition: background 0.3s, color 0.3s;}
 .header{background:var(--primary);color:#fff;padding:15px;text-align:center;font-size:18px;font-weight:bold;border-radius:0 0 15px 15px;margin-bottom:15px; display:flex; justify-content:space-between; align-items:center;}
 .header a {color: white; text-decoration: none; font-size: 20px;}
 .card{background:var(--card);padding:15px;border-radius:12px;box-shadow:0 2px 5px #0001;margin:0 10px 15px 10px;border:1px solid var(--border); color: var(--text);}
 h3{color:var(--text);margin-top:0;text-align:center;font-size:16px}
 input,select,textarea{width:100%;padding:10px;margin:5px 0;box-sizing:border-box;border:1px solid var(--border);border-radius:8px;font-size:14px;background:var(--input-bg);color:var(--text); font-family: inherit;}
 .row{display:flex;gap:5px}.row input,.row select{margin:0}
 button{background:var(--btn);color:#fff;border:0;padding:12px;width:100%;border-radius:8px;font-size:16px;font-weight:bold;margin-top:10px;cursor:pointer}
 button:active{opacity: 0.8;}
 button.del{background:#e53935;padding:6px;width:auto;font-size:12px;margin:0}
 button.export{background:#1565c0;margin-bottom:10px;width: 32%}
 button.import{background:#f57c00;margin-bottom:10px;width: 32%}
 button.orden{background:#8e24aa;margin-bottom:10px;width: 32%}
 button.pdf{background:#d32f2f;margin-bottom:10px;width: 100%}
 button.foto{background:#ff9800;margin-bottom:0px;} 
 button.tab-resumen{background:var(--border);color:var(--text); padding:10px; font-size:13px; width:auto; flex:1; margin:0 2px;}
 button.tab-resumen.active{background:var(--primary); color:#fff;}
 button.peligro{background:#b71c1c; margin-top:20px;}
 
 .table-container{width: 100%; overflow-x: auto; margin-top:10px; border-radius:8px;}
 table{width:100%; border-collapse:collapse; background:var(--card); font-size:12px; min-width: 900px; color: var(--text);}
 th,td{padding:8px;border-bottom:1px solid var(--border);text-align:left; white-space: nowrap;}
 th{background:var(--table-th);color:#fff}
 
 .totales{background:var(--border);padding:10px;border-radius:8px;margin-top:15px;font-weight:bold;font-size:13px;color:var(--text)}
 .tabs{display:flex;background:var(--card);position:fixed;bottom:0;width:100%;box-shadow:0 -2px 10px #0001;z-index:100}
 .tab{flex:1;text-align:center;padding:10px 2px;text-decoration:none;color:var(--text);font-weight:bold;font-size:11px;border-top:3px solid transparent}
 .tab.active{color:var(--btn);border-top:3px solid var(--btn);background:var(--bg)}
 .chat-box { background:var(--input-bg); border-radius:12px; padding:10px; height: 350px; overflow-y:auto; margin-bottom:10px; display:flex; flex-direction:column; gap:8px; border:1px solid var(--border);}
 .msg-user { background:var(--border); padding:8px 12px; border-radius:12px; align-self:flex-end; max-width:80%; font-size:14px;}
 .msg-ia { background:var(--card); border: 1px solid var(--border); padding:8px 12px; border-radius:12px; align-self:flex-start; max-width:80%; font-size:14px;}
 
 .nav-resumen{display:flex; justify-content:space-between; margin-bottom: 15px;}
 .global-card{background:var(--primary); color: white; padding: 15px; border-radius: 12px; text-align: center; margin-bottom: 15px;}
 .global-card h4{margin: 0; font-size: 14px; opacity: 0.8;}
 .global-card .monto{font-size: 24px; font-weight: bold; margin-top: 5px;}
 .resumen-box {background:var(--input-bg); padding: 15px; border-radius: 12px; border: 1px solid var(--border); margin-bottom: 15px;}
 .resumen-box h4 {margin: 0 0 10px 0; color: var(--text); font-size: 16px; border-bottom: 2px solid var(--btn); padding-bottom: 5px;}
 .resumen-line {display: flex; justify-content: space-between; font-size: 14px; margin-bottom: 8px;}
 .resumen-line.total {border-top: 1px solid var(--border); padding-top: 8px; margin-top: 8px; font-weight: bold; font-size: 16px;}
 .text-red {color: #ff5252;}
 .text-green {color: #69f0ae;}
 
 .login-box {max-width: 300px; margin: 50px auto; padding: 20px; background: var(--card); border-radius: 12px; box-shadow: 0 4px 10px #0002;}
</style>
<script>
  document.addEventListener("DOMContentLoaded", function() {
    if(localStorage.getItem("theme") === "dark") { document.body.classList.add("dark-mode"); }
  });
  function toggleTheme() {
    document.body.classList.toggle("dark-mode");
    if(document.body.classList.contains("dark-mode")) { localStorage.setItem("theme", "dark"); } 
    else { localStorage.setItem("theme", "light"); }
  }
</script>
</head><body>
"""

HTML_HEADER = """
<div class="header">
  <span>🌱 Mi Libro</span>
  <div>
    <a href="#" onclick="toggleTheme()" title="Modo Oscuro/Claro">🌓</a>
    <a href="/perfil" style="margin-left: 15px;">👤</a>
  </div>
</div>
"""

HTML_FOOTER = "</body></html>"

HTML_LOGIN = HTML_BASE + """
<div class="login-box">
  <h3>Iniciar Sesión</h3>
  {% if error %}<p style="color:red; text-align:center;">{{ error }}</p>{% endif %}
  <form method="POST" action="/login">
    <input type="text" name="username" placeholder="Usuario" required>
    <input type="password" name="password" placeholder="Contraseña" required>
    <button type="submit">Entrar</button>
  </form>
  <p style="text-align:center; font-size:12px; margin-top:15px;">¿No tienes cuenta? <a href="/register" style="color:var(--btn)">Regístrate aquí</a></p>
</div>
""" + HTML_FOOTER

HTML_REGISTER = HTML_BASE + """
<div class="login-box">
  <h3>Crear Cuenta</h3>
  {% if error %}<p style="color:red; text-align:center;">{{ error }}</p>{% endif %}
  <form method="POST" action="/register">
    <input type="text" name="username" placeholder="Elige un Usuario" required>
    <input type="password" name="password" placeholder="Elige una Contraseña" required>
    <button type="submit">Registrarse</button>
  </form>
  <p style="text-align:center; font-size:12px; margin-top:15px;">¿Ya tienes cuenta? <a href="/login" style="color:var(--btn)">Inicia sesión</a></p>
</div>
""" + HTML_FOOTER

HTML_PERFIL = HTML_BASE + HTML_HEADER + """
<div class="card">
  <h3>Mi Perfil</h3>
  <p style="text-align:center; font-size:18px;">Conectado como: <b>{{ user }}</b></p>
  <a href="/logout"><button style="background:#e53935; margin-top:20px;">Cerrar Sesión</button></a>
</div>
<div class="tabs">
 <a href="/?vista=compras" class="tab">📗 Volver al Libro</a>
</div>
""" + HTML_FOOTER

HTML_APP = HTML_BASE + HTML_HEADER + """
{% if vista == 'compras' or vista == 'ventas' or vista == 'todos' %}
{% if vista != 'todos' %}
<div class="card">
 <h3>Registrar {{ 'Compra' if vista == 'compras' else 'Venta' }}</h3>
 <form action="/procesar_foto" method="post" enctype="multipart/form-data" style="margin-bottom:15px;">
   <input type="hidden" name="vista_actual" value="{{ vista }}">
   <input type="file" name="foto" accept="image/*" capture="environment" style="display:none" id="camara" onchange="this.form.submit()">
   <label for="camara"><button type="button" class="foto" onclick="document.getElementById('camara').click()">📸 Tomar Foto</button></label>
   <p style="font-size: 12px; color: #d32f2f; text-align: center; margin-top: 5px; font-weight: bold; margin-bottom: 0;">⚠️ Revisar antes de descargar</p>
 </form>
 {% if mensaje_foto %}
 <div style="background:#fff3e0; color:#e65100; padding:10px; border-radius:8px; margin-bottom:10px; font-size:13px; text-align:center;">{{ mensaje_foto }}</div>
 {% endif %}
 <form method="post" action="/add">
  <input type="hidden" name="tipo" value="{{ 'Compra' if vista == 'compras' else 'Venta' }}">
  <div class="row">
    <select name="tipo_doc" id="tipo_doc" onchange="calcular()">
      {% if vista == 'compras' %}
        <option value="Factura Electrónica">Factura Electrónica</option>
        <option value="Factura No Electrónica">Factura No Electrónica</option>
        <option value="Boleta">Boleta</option>
        <option value="Nota de Crédito">Nota de Crédito</option>
        <option value="Nota de Débito">Nota de Débito</option>
        <option value="Factura de Compra">Factura de Compra</option>
        <option value="Importación">Importación</option>
        <option value="Otro">Otro</option>
        <option value="Factura Maquinaria">Factura Maquinaria</option>
      {% else %}
        <option value="Factura Electrónica">Factura Electrónica</option>
        <option value="Boleta Electrónica">Boleta Electrónica</option>
        <option value="Boleta">Boleta</option>
        <option value="Nota de Crédito">Nota de Crédito</option>
        <option value="Nota de Débito">Nota de Débito</option>
        <option value="Otro">Otro</option>
      {% endif %}
    </select>
    <input name="fecha" type="date" value="{{ datos_foto.fecha if datos_foto else hoy }}" required>
  </div>
  <div class="row">
    <input name="folio" placeholder="Folio" value="{{ datos_foto.folio if datos_foto else '' }}" required>
    <input name="rut" id="rut" placeholder="RUT" value="{{ datos_foto.rut if datos_foto else '' }}" required>
  </div>
  <input name="nombre" id="nombre" placeholder="Nombre {{ 'Proveedor' if vista == 'compras' else 'Cliente' }}" value="{{ datos_foto.nombre if datos_foto else '' }}" required onblur="buscarRut()">
  <div class="row">
    <input name="neto" id="neto" placeholder="Neto" type="number" step="1" value="{{ datos_foto.neto if datos_foto else 0 }}" required oninput="calcular()">
    <input name="exento" id="exento" placeholder="Exento" type="number" step="1" value="{{ datos_foto.exento if datos_foto else 0 }}" oninput="calcular()">
  </div>
  <div class="row">
    <input name="iva" id="iva" placeholder="IVA" type="number" step="1" value="{{ datos_foto.iva if datos_foto else 0 }}" readonly>
    <input name="total" id="total" placeholder="Total" type="number" step="1" value="{{ datos_foto.total if datos_foto else 0 }}" readonly>
  </div>
  <select name="pago"><option>Contado</option><option>Crédito</option><option>Transferencia</option></select>
  <textarea name="nota" placeholder="Notas" rows="2">{{ datos_foto.nota if datos_foto else '' }}</textarea>
  <button>💾 Registrar</button>
 </form>
</div>
{% endif %}

<div class="card">
 <h3>{{ 'Documentos Registrados' if vista == 'todos' else 'Libro de ' + ('Compras' if vista == 'compras' else 'Ventas') }}</h3>
 
 {% if vista != 'todos' %}
 <div style="overflow:hidden; margin-bottom:10px;">
   <a href="/exportar_excel?vista={{vista}}"><button class="export">📥 Excel</button></a>
   <a href="#" onclick="document.getElementById('import_excel').click()"><button class="import">📤 Importar</button></a>
   <a href="/?vista=orden"><button class="orden">⚙️ Orden</button></a>
   <form action="/importar_excel" method="POST" enctype="multipart/form-data" style="display:none;" id="form_import">
     <input type="file" name="excel_file" id="import_excel" accept=".xlsx" onchange="document.getElementById('form_import').submit()">
     <input type="hidden" name="vista_actual" value="{{vista}}">
   </form>
 </div>
 <div><a href="/exportar_pdf?vista={{vista}}"><button class="pdf">📄 Descargar PDF</button></a></div>
 {% endif %}

 <div class="table-container">
     <table>
      <tr>
        <th>Tipo</th><th>Tipo Doc.</th><th>Fecha</th><th>Folio</th><th>RUT</th><th>Nombre</th>
        <th>Neto</th><th>Exento</th><th>IVA</th><th>Total</th><th>{{ 'Pago' if vista == 'compras' else 'Cobro' }}</th><th>Nota</th><th>Acción</th>
      </tr>
      {% for d in docs %}
      <tr>
        <td>{{ d[1] }}</td><td>{{ d[2] }}</td><td>{{ d[3] }}</td><td>{{ d[4] }}</td><td>{{ d[5] }}</td><td>{{ d[6] }}</td>
        <td>${{ '{:,}'.format(d[7]|float|int).replace(',','.') }}</td><td>${{ '{:,}'.format(d[8]|float|int).replace(',','.') }}</td>
        <td>${{ '{:,}'.format(d[9]|float|int).replace(',','.') }}</td><td>${{ '{:,}'.format(d[10]|float|int).replace(',','.') }}</td>
        <td>{{ d[11] }}</td><td style="color: #d84315; font-style: italic;">{{ d[12] }}</td>
        <td><a href="/del/{{d[0]}}?vista={{vista}}"><button class="del">X</button></a></td>
      </tr>
      {% endfor %}
     </table>
 </div>
 {% set ns = namespace(neto=0.0, exento=0.0, iva=0.0, total=0.0) %}
 {% for d in docs %}
    {% set ns.neto = ns.neto + (d[7]|float) %}{% set ns.exento = ns.exento + (d[8]|float) %}
    {% set ns.iva = ns.iva + (d[9]|float) %}{% set ns.total = ns.total + (d[10]|float) %}
 {% endfor %}
 <div class="totales">
  Neto: ${{'{:,}'.format(ns.neto|int).replace(',','.')}} | Exento: ${{'{:,}'.format(ns.exento|int).replace(',','.')}} | 
  IVA: ${{'{:,}'.format(ns.iva|int).replace(',','.')}} | Total: ${{'{:,}'.format(ns.total|int).replace(',','.')}}
 </div>
</div>

{% elif vista == 'resumen' %}
<div class="card">
 <div class="global-card">
   <h4>Total Neto General</h4><div class="monto">${{'{:,}'.format(global_neto|int).replace(',','.')}}</div>
   <div style="font-size:12px; margin-top:5px;">Documentos registrados: {{ global_total_docs }}</div>
 </div>
 <div style="text-align:center; margin-bottom: 15px;">
   <a href="/exportar_ordenado_excel"><button style="background:#1565c0; width: 100%; margin-bottom:5px;">📥 Resumen General (Excel)</button></a>
   <a href="/exportar_pdf?vista=resumen"><button style="background:#d32f2f; width: 100%;">📄 Resumen General (PDF)</button></a>
 </div>
 <div class="nav-resumen">
   <a href="/?vista=resumen&tipo=compras" style="flex:1; text-decoration:none;"><button class="tab-resumen {{ 'active' if tipo_resumen == 'compras' else '' }}">📗 Compras</button></a>
   <a href="/?vista=resumen&tipo=ventas" style="flex:1; text-decoration:none;"><button class="tab-resumen {{ 'active' if tipo_resumen == 'ventas' else '' }}">📘 Ventas</button></a>
   <a href="/?vista=resumen&tipo=general" style="flex:1; text-decoration:none;"><button class="tab-resumen {{ 'active' if tipo_resumen == 'general' else '' }}">📊 General</button></a>
 </div>
 {% if tipo_resumen == 'compras' %}
   <div class="resumen-box"><h4>Resumen de Compras</h4>
     <div class="resumen-line"><span>Neto:</span> <strong>${{'{:,}'.format(tot_c.neto|int).replace(',','.')}}</strong></div>
     <div class="resumen-line"><span>IVA Crédito:</span> <strong>${{'{:,}'.format(tot_c.iva|int).replace(',','.')}}</strong></div>
     <div class="resumen-line total"><span>TOTAL COMPRAS:</span> <span>${{'{:,}'.format(tot_c.total|int).replace(',','.')}}</span></div>
   </div>
 {% elif tipo_resumen == 'ventas' %}
   <div class="resumen-box"><h4>Resumen de Ventas</h4>
     <div class="resumen-line"><span>Neto:</span> <strong>${{'{:,}'.format(tot_v.neto|int).replace(',','.')}}</strong></div>
     <div class="resumen-line"><span>IVA Débito:</span> <strong>${{'{:,}'.format(tot_v.iva|int).replace(',','.')}}</strong></div>
     <div class="resumen-line total"><span>TOTAL VENTAS:</span> <span>${{'{:,}'.format(tot_v.total|int).replace(',','.')}}</span></div>
   </div>
 {% elif tipo_resumen == 'general' %}
   <div class="resumen-box"><h4>F29 - Resumen General</h4>
     <div class="resumen-line" style="color:#1565c0;"><span>IVA Crédito:</span> <strong>${{'{:,}'.format(tot_c.iva|int).replace(',','.')}}</strong></div>
     <div class="resumen-line" style="color:#f57c00;"><span>IVA Débito:</span> <strong>${{'{:,}'.format(tot_v.iva|int).replace(',','.')}}</strong></div>
     {% set dif = tot_v.iva - tot_c.iva %}
     <div class="resumen-line total" style="font-size: 16px; margin-top: 15px;">
        <span>Resultado IVA:</span> 
        {% if dif > 0 %}<span class="text-red">${{'{:,}'.format(dif|int).replace(',','.')}} (A Pagar)</span>
        {% elif dif < 0 %}<span class="text-green">${{'{:,}'.format((dif * -1)|int).replace(',','.')}} (A Favor)</span>
        {% else %}<span>$0</span>{% endif %}
     </div>
   </div>
   <div class="resumen-box"><h4>Flujo Total</h4>
      <div class="resumen-line"><span>Total Compras:</span> <strong>${{'{:,}'.format(tot_c.total|int).replace(',','.')}}</strong></div>
      <div class="resumen-line"><span>Total Ventas:</span> <strong>${{'{:,}'.format(tot_v.total|int).replace(',','.')}}</strong></div>
      <div class="resumen-line total"><span>Movimiento Total:</span> <span>${{'{:,}'.format((tot_c.total + tot_v.total)|int).replace(',','.')}}</span></div>
   </div>
 {% endif %}
 <div style="text-align:center; margin-top: 30px; padding: 0 10px;">
    <form action="/nuevo_libro" method="POST" onsubmit="return confirm('¿Empezar un libro nuevo? Se borrarán TUS datos actuales.');">
        <button type="submit" class="peligro">🔄 Empezar Libro Nuevo</button>
    </form>
 </div>
</div>

{% elif vista == 'orden' %}
<div class="card">
  <h3>⚙️ Ordenar libro</h3>
  <p style="font-size:13px; text-align:center; color:var(--text);">
    Genera un Excel con formato del liceo:<br><br>
    1° Compras (Factura+Maquinaria, Nota Déb, Nota Créd)<br>2° Ventas<br><br>
    <i>* Orden alfabético (Proveedor), Fecha, Folio.</i>
  </p>
  <a href="/exportar_ordenado_excel"><button style="background:#8e24aa;">📥 Descargar Excel Ordenado</button></a>
</div>

{% elif vista == 'ia' %}
<div class="card" style="height: 500px; display:flex; flex-direction:column;">
 <h3>🤖 Asistente IA</h3>
 <div class="chat-box" id="chat-box"><div class="msg-ia">¡Hola! Pregúntame sobre contabilidad o el IVA.</div></div>
 <form id="form-ia" style="margin-top:auto;">
   <div class="row"><input id="pregunta" placeholder="Tu duda..." required><button type="submit" style="width:auto; padding:10px 15px; margin:0;">Enviar</button></div>
 </form>
</div>
<script>
document.getElementById('form-ia').onsubmit = async (e) => {
    e.preventDefault(); let input = document.getElementById('pregunta'); let texto = input.value; let chat = document.getElementById('chat-box');
    chat.innerHTML += `<div class="msg-user">${texto}</div>`; input.value = ''; chat.scrollTop = chat.scrollHeight;
    let idTemp = "temp-" + Date.now(); chat.innerHTML += `<div class="msg-ia" id="${idTemp}">Pensando...</div>`; chat.scrollTop = chat.scrollHeight;
    try {
        let res = await fetch('/api/ia', { method: 'POST', headers: {'Content-Type': 'application/json'}, body: JSON.stringify({pregunta: texto}) });
        let data = await res.json(); document.getElementById(idTemp).innerText = data.respuesta;
    } catch(err) { document.getElementById(idTemp).innerText = "Error."; }
    chat.scrollTop = chat.scrollHeight;
};
</script>
{% endif %}

<div class="tabs">
 <a href="/?vista=compras" class="tab {{ 'active' if vista=='compras' }}">📗 Compras</a>
 <a href="/?vista=ventas" class="tab {{ 'active' if vista=='ventas' }}">📘 Ventas</a>
 <a href="/?vista=todos" class="tab {{ 'active' if vista=='todos' }}">📑 Todos</a>
 <a href="/?vista=resumen&tipo=general" class="tab {{ 'active' if vista=='resumen' }}">📊 Resumen</a>
 <a href="/?vista=ia" class="tab {{ 'active' if vista=='ia' }}">🤖 IA</a>
</div>

<script>
function calcular() {
    let neto = parseFloat(document.getElementById("neto").value) || 0;
    let exento = parseFloat(document.getElementById("exento").value) || 0;
    let tipo_doc = document.getElementById("tipo_doc").value;
    let iva = 0; if (!tipo_doc.includes("Exenta")) { iva = Math.round(neto * 0.19); }
    let total = neto + exento + iva;
    if (tipo_doc === "Nota de Crédito") { iva = -iva; total = -total; }
    document.getElementById("iva").value = iva; document.getElementById("total").value = total;
}
async function buscarRut() {
    let nombreInput = document.getElementById("nombre").value.trim(); let rutInput = document.getElementById("rut");
    if (rutInput.value.trim() !== "" || nombreInput.length < 3) return; 
    try {
        let res = await fetch(`/api/get_rut?nombre=${encodeURIComponent(nombreInput)}`); let data = await res.json();
        if (data.rut) { rutInput.value = data.rut; }
    } catch(e) {}
}
</script>
""" + HTML_FOOTER

# ==========================================
# RUTAS DE AUTENTICACIÓN
# ==========================================
@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        user = request.form['username']; pw = request.form['password']
        con = get_db(); cur = con.cursor()
        cur.execute("SELECT id, password FROM usuarios WHERE username=%s", (user,))
        row = cur.fetchone(); cur.close(); con.close()
        if row and check_password_hash(row[1], pw):
            session['user_id'] = row[0]; session['username'] = user
            return redirect('/')
        return render_template_string(HTML_LOGIN, error="Usuario o contraseña incorrectos")
    return render_template_string(HTML_LOGIN)

@app.route('/register', methods=['GET', 'POST'])
def register():
    if request.method == 'POST':
        user = request.form['username']; pw = request.form['password']
        hashed_pw = generate_password_hash(pw)
        try:
            con = get_db(); cur = con.cursor()
            cur.execute("INSERT INTO usuarios (username, password) VALUES (%s, %s)", (user, hashed_pw))
            con.commit(); cur.close(); con.close()
            return redirect('/login')
        except psycopg2.IntegrityError:
            return render_template_string(HTML_REGISTER, error="Ese usuario ya existe")
    return render_template_string(HTML_REGISTER)

@app.route('/logout')
def logout():
    session.clear()
    return redirect('/login')

@app.route('/perfil')
@login_required
def perfil():
    return render_template_string(HTML_PERFIL, user=session['username'])

# ==========================================
# RUTAS PRINCIPALES DE LA APP
# ==========================================
@app.route("/")
@login_required
def home():
    vista = request.args.get("vista", "compras")
    tipo_resumen = request.args.get("tipo", "general") 
    uid = session['user_id']
    
    con = get_db(); cur = con.cursor()
    
    if vista == "resumen":
        cur.execute("SELECT count(*), sum(cast(neto as real)) FROM docs WHERE user_id=%s", (uid,))
        row = cur.fetchone(); global_docs = row[0] or 0; global_neto = row[1] or 0.0
        cur.execute("SELECT sum(cast(neto as real)), sum(cast(exento as real)), sum(cast(iva as real)), sum(cast(total as real)) FROM docs WHERE tipo='Compra' AND user_id=%s", (uid,))
        rc = cur.fetchone(); tot_c = {"neto": rc[0] or 0, "exento": rc[1] or 0, "iva": rc[2] or 0, "total": rc[3] or 0}
        cur.execute("SELECT sum(cast(neto as real)), sum(cast(exento as real)), sum(cast(iva as real)), sum(cast(total as real)) FROM docs WHERE tipo='Venta' AND user_id=%s", (uid,))
        rv = cur.fetchone(); tot_v = {"neto": rv[0] or 0, "exento": rv[1] or 0, "iva": rv[2] or 0, "total": rv[3] or 0}
        cur.close(); con.close()
        return render_template_string(HTML_APP, vista=vista, tipo_resumen=tipo_resumen, global_docs=global_docs, global_neto=global_neto, tot_c=tot_c, tot_v=tot_v)

    if vista == "ia" or vista == "orden":
        cur.close(); con.close()
        return render_template_string(HTML_APP, vista=vista, hoy="")

    if vista == "todos":
        cur.execute("SELECT id, tipo, tipo_doc, fecha, folio, rut, nombre, neto, exento, iva, total, pago, nota FROM docs WHERE user_id=%s ORDER BY fecha DESC", (uid,))
        docs = cur.fetchall(); cur.close(); con.close()
        return render_template_string(HTML_APP, docs=docs, vista=vista, hoy=datetime.now().strftime("%Y-%m-%d"))

    tipo = "Compra" if vista == "compras" else "Venta"
    cur.execute("SELECT id, tipo, tipo_doc, fecha, folio, rut, nombre, neto, exento, iva, total, pago, nota FROM docs WHERE tipo=%s AND user_id=%s ORDER BY fecha DESC", (tipo, uid))
    docs = cur.fetchall(); cur.close(); con.close()
    return render_template_string(HTML_APP, docs=docs, vista=vista, hoy=datetime.now().strftime("%Y-%m-%d"))

@app.route("/add", methods=["POST"])
@login_required
def add():
    f = request.form
    uid = session['user_id']
    neto, exento = float(f["neto"] or 0), float(f["exento"] or 0)
    tipo_doc, nota = f["tipo_doc"], f.get("nota", "")
    
    if tipo_doc == "Factura Maquinaria":
        tipo_doc = "Factura Electrónica"
        if nota:
            nota = "Maquinaria - " + nota
        else:
            nota = "Maquinaria"

    iva = 0 if "Exenta" in tipo_doc else round(neto * 0.19)
    total = neto + exento + iva
    if tipo_doc == "Nota de Crédito": neto, exento, iva, total = -abs(neto), -abs(exento), -abs(iva), -abs(total)
    
    con = get_db(); cur = con.cursor()
    cur.execute("""INSERT INTO docs(user_id,tipo,tipo_doc,fecha,folio,rut,nombre,neto,exento,iva,total,pago,nota) 
                   VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)""", (uid, f["tipo"], tipo_doc, f["fecha"], f["folio"], f["rut"], f["nombre"], neto, exento, iva, total, f["pago"], nota))
    con.commit(); cur.close(); con.close()
    return redirect(f"/?vista={'compras' if f['tipo'] == 'Compra' else 'ventas'}")

@app.route("/del/<int:id>")
@login_required
def delete(id):
    uid = session['user_id']
    con = get_db(); cur = con.cursor()
    cur.execute("DELETE FROM docs WHERE id=%s AND user_id=%s", (id, uid))
    con.commit(); cur.close(); con.close()
    return redirect(f"/?vista={request.args.get('vista', 'compras')}")

@app.route("/nuevo_libro", methods=["POST"])
@login_required
def nuevo_libro():
    uid = session['user_id']
    con = get_db(); cur = con.cursor()
    cur.execute("DELETE FROM docs WHERE user_id=%s", (uid,))
    con.commit(); cur.close(); con.close()
    return redirect("/?vista=compras")

# ==========================================
# IMPORTAR / EXPORTAR
# ==========================================
@app.route("/importar_excel", methods=["POST"])
@login_required
def importar_excel():
    uid = session['user_id']
    vista_actual = request.form.get("vista_actual", "compras")
    if 'excel_file' not in request.files: return redirect(f"/?vista={vista_actual}")
    file = request.files['excel_file']
    if file.filename == '': return redirect(f"/?vista={vista_actual}")

    try:
        wb = load_workbook(file)
        ws = wb.active
        con = get_db(); cur = con.cursor()
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row[0]: continue
            tipo = "Compra" if vista_actual == "compras" else "Venta"
            cur.execute("""INSERT INTO docs(user_id,tipo,tipo_doc,fecha,folio,rut,nombre,neto,exento,iva,total,pago,nota) 
                           VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)""", 
                        (uid, tipo, row[1], row[0], row[2], row[3], row[4], row[5], row[6], row[7], row[8], row[9], row[10]))
        con.commit(); cur.close(); con.close()
    except Exception as e:
        print("Error importando:", e)

    return redirect(f"/?vista={vista_actual}")

@app.route("/exportar_excel")
@login_required
def exportar_excel():
    uid = session['user_id']
    vista = request.args.get("vista", "compras")
    tipo = "Compra" if vista == "compras" else "Venta"
    con = get_db(); cur = con.cursor()
    cur.execute("SELECT fecha, tipo_doc, folio, rut, nombre, neto, exento, iva, total, pago, nota FROM docs WHERE tipo=%s AND user_id=%s ORDER BY fecha", (tipo, uid))
    docs = cur.fetchall(); cur.close(); con.close()

    wb = Workbook(); ws = wb.active; ws.title = f"Libro {vista.capitalize()}"
    headers = ["Fecha", "TipoDoc", "Folio", "RUT", "Nombre", "Neto", "Exento", "IVA", "Total", "Pago/Cobro", "Nota"]
    ws.append(headers)
    for d in docs: ws.append(list(d))
    
    for col in ws.iter_cols(min_row=1, max_row=1):
        for cell in col:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="2E7D32", end_color="2E7D32", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
    for column_cells in ws.columns:
        max_length = 0
        column_letter = column_cells[0].column_letter
        for cell in column_cells:
            try:
                if len(str(cell.value)) > max_length: max_length = len(str(cell.value))
            except: pass
        ws.column_dimensions[column_letter].width = max_length + 2
        
    thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row: cell.border = thin_border
    
    mem = io.BytesIO(); wb.save(mem); mem.seek(0)
    return send_file(mem, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', as_attachment=True, download_name=f"Libro_{vista}.xlsx")

@app.route("/exportar_ordenado_excel")
@login_required
def exportar_ordenado_excel():
    uid = session['user_id']
    con = get_db(); cur = con.cursor()
    cur.execute("SELECT tipo, tipo_doc, fecha, folio, rut, nombre, neto, exento, iva, total, pago, nota FROM docs WHERE user_id=%s", (uid,))
    all_docs = cur.fetchall(); cur.close(); con.close()
    
    def sort_func(d): return (str(d[5]).lower(), str(d[2]), str(d[3]))
    compras = [d for d in all_docs if d[0] == 'Compra']
    c_fact = sorted([d for d in compras if "Factura" in d[1]], key=sort_func)
    c_nota_deb = sorted([d for d in compras if d[1] == "Nota de Débito"], key=sort_func)
    c_nota_cred = sorted([d for d in compras if d[1] == "Nota de Crédito"], key=sort_func)
    c_otros = sorted([d for d in compras if d not in c_fact and d not in c_nota_deb and d not in c_nota_cred], key=sort_func)

    ventas = [d for d in all_docs if d[0] == 'Venta']
    v_fact = sorted([d for d in ventas if d[1] == "Factura Electrónica"], key=sort_func)
    v_bole_elec = sorted([d for d in ventas if d[1] == "Boleta Electrónica"], key=sort_func)
    v_bole = sorted([d for d in ventas if d[1] == "Boleta"], key=sort_func)
    v_nota_cred = sorted([d for d in ventas if d[1] == "Nota de Crédito"], key=sort_func)
    v_nota_deb = sorted([d for d in ventas if d[1] == "Nota de Débito"], key=sort_func)
    v_otros = sorted([d for d in ventas if d not in v_fact and d not in v_bole_elec and d not in v_bole and d not in v_nota_cred and d not in v_nota_deb], key=sort_func)

    def sum_g(g):
        return (round(sum(float(d[6] or 0) for d in g), 2), round(sum(float(d[7] or 0) for d in g), 2), round(sum(float(d[8] or 0) for d in g), 2), round(sum(float(d[9] or 0) for d in g), 2))

    s_c_fact, s_c_deb, s_c_cred, s_c_otros = sum_g(c_fact), sum_g(c_nota_deb), sum_g(c_nota_cred), sum_g(c_otros)
    s_c_tot = (round(s_c_fact[0] + s_c_deb[0] - abs(s_c_cred[0]) + s_c_otros[0], 2), round(s_c_fact[1] + s_c_deb[1] - abs(s_c_cred[1]) + s_c_otros[1], 2), round(s_c_fact[2] + s_c_deb[2] - abs(s_c_cred[2]) + s_c_otros[2], 2), round(s_c_fact[3] + s_c_deb[3] - abs(s_c_cred[3]) + s_c_otros[3], 2))
    s_v_tot = sum_g(ventas)

    wb = Workbook(); ws = wb.active; ws.title = "Resumen Ordenado"
    ws.append(["Libro", "TipoDoc", "Fecha", "Folio", "RUT", "Nombre", "Neto", "Exento", "IVA", "Total", "Pago/Cobro", "Nota"])
    def write_rows(rows):
        for d in rows: ws.append([d[0], d[1], d[2], d[3], d[4], d[5], d[6], d[7], d[8], d[9], d[10], d[11]])
    write_rows(c_fact)
    ws.append(["TOTAL FACTURA Y MAQUINARIA", "", "", "", "", "", s_c_fact[0], s_c_fact[1], s_c_fact[2], s_c_fact[3], "", ""])
    ws.append([]); write_rows(c_nota_deb)
    ws.append(["TOTAL NOTA DE DÉBITO", "", "", "", "", "", s_c_deb[0], s_c_deb[1], s_c_deb[2], s_c_deb[3], "", ""])
    ws.append([]); write_rows(c_nota_cred)
    ws.append(["TOTAL NOTA DE CRÉDITO", "", "", "", "", "", s_c_cred[0], s_c_cred[1], s_c_cred[2], s_c_cred[3], "", ""])
    ws.append([])
    if c_otros:
        write_rows(c_otros); ws.append(["TOTAL OTROS", "", "", "", "", "", s_c_otros[0], s_c_otros[1], s_c_otros[2], s_c_otros[3], "", ""]); ws.append([])
    ws.append(["TOTAL LIBRO COMPRAS", "", "", "", "", "", s_c_tot[0], s_c_tot[1], s_c_tot[2], s_c_tot[3], "", ""])
    ws.append([]); ws.append([])
    write_rows(v_fact); write_rows(v_bole_elec); write_rows(v_bole); write_rows(v_nota_cred); write_rows(v_nota_deb); write_rows(v_otros)
    ws.append([])
    ws.append(["TOTAL LIBRO VENTAS", "", "", "", "", "", s_v_tot[0], s_v_tot[1], s_v_tot[2], s_v_tot[3], "", ""])
    
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        if row[0].value and str(row[0].value).startswith("TOTAL"):
            for cell in row:
                cell.font = Font(bold=True)
                cell.fill = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")

    mem = io.BytesIO(); wb.save(mem); mem.seek(0)
    return send_file(mem, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', as_attachment=True, download_name="Resumen_Liceo.xlsx")

@app.route("/exportar_pdf")
@login_required
def exportar_pdf():
    uid = session['user_id']
    vista = request.args.get("vista", "resumen")
    con = get_db(); cur = con.cursor()
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter)
    elementos = [Paragraph(f"<b>Reporte Contable - {vista.capitalize()}</b>", getSampleStyleSheet()['Title']), Spacer(1, 12)]
    
    if vista == "resumen":
        cur.execute("SELECT sum(cast(iva as real)) FROM docs WHERE tipo='Compra' AND user_id=%s", (uid,)); iva_c = cur.fetchone()[0] or 0
        cur.execute("SELECT sum(cast(iva as real)) FROM docs WHERE tipo='Venta' AND user_id=%s", (uid,)); iva_v = cur.fetchone()[0] or 0
        dif = iva_v - iva_c
        data = [["Concepto", "Monto"], ["IVA Crédito", f"${iva_c:,.0f}"], ["IVA Débito", f"${iva_v:,.0f}"], ["Resultado IVA", f"${dif:,.0f}"]]
        t = Table(data, colWidths=[300, 150])
    else:
        tipo = "Compra" if vista == "compras" else "Venta"
        cur.execute("SELECT fecha, tipo_doc, nombre, total FROM docs WHERE tipo=%s AND user_id=%s ORDER BY fecha", (tipo, uid))
        data = [["Fecha", "Documento", "Nombre", "Total"]] + [[r[0], r[1], r[2], f"${float(r[3] or 0):,.0f}"] for r in cur.fetchall()]
        t = Table(data, colWidths=[80, 120, 200, 80])
        
    t.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#2E7D32")), ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke), ('GRID', (0, 0), (-1, -1), 1, colors.HexColor("#c8e6c9"))]))
    elementos.append(t); cur.close(); con.close(); doc.build(elementos); buffer.seek(0)
    return send_file(buffer, mimetype='application/pdf', as_attachment=True, download_name=f"Reporte_{vista}.pdf")

# ==========================================
# IA (FOTO Y CHAT) - CON SEGURIDAD DE CLAVE
# ==========================================
@app.route("/procesar_foto", methods=["POST"])
@login_required
def procesar_foto():
    vista_actual = request.form.get("vista_actual", "compras")
    file = request.files.get('foto')
    if not file or file.filename == '': return redirect(f"/?vista={vista_actual}")
    
    if not GROQ_API_KEY: return "Falta configurar la API KEY en el servidor", 500

    try:
        img_bytes = file.read(); base64_image = base64.b64encode(img_bytes).decode('utf-8')
        client = Groq(api_key=GROQ_API_KEY)
        prompt = "Eres un asistente contable. Devuelve ESTRICTAMENTE una lista JSON: [{'rut':'...', 'folio':'...', 'nombre':'...', 'neto':0, 'iva':0, 'total':0, 'fecha':'...'}]"
        
        completion = client.chat.completions.create(model="openai/gpt-oss-20b", messages=[{"role": "user", "content": [{"type": "text", "text": prompt}, {"type": "image_url", "image_url": {"url": f"data:image/jpeg;base64,{base64_image}"}}]}], temperature=0.1)
        datos_lista = json.loads(completion.choices[0].message.content.replace('```json', '').replace('```', '').strip())
        if isinstance(datos_lista, dict): datos_lista = [datos_lista]
            
        tipo = "Compra" if vista_actual == "compras" else "Venta"
        uid = session['user_id']
        con = get_db(); cur = con.cursor(); guardados = 0
        for datos in datos_lista:
            if not datos.get("total"): continue 
            neto = float(datos.get("neto") or 0); iva = float(datos.get("iva") or (neto * 0.19)); total = float(datos.get("total") or (neto + iva))
            cur.execute("""INSERT INTO docs(user_id,tipo,tipo_doc,fecha,folio,rut,nombre,neto,exento,iva,total,pago,nota) VALUES(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)""", 
                        (uid, tipo, "Factura Electrónica", datos.get("fecha") or datetime.now().strftime("%Y-%m-%d"), datos.get("folio") or "S/N", datos.get("rut") or "S/R", datos.get("nombre") or "Desc", neto, 0, iva, total, "Contado", ""))
            guardados += 1
        con.commit(); cur.close(); con.close()
        return redirect(f"/?vista={vista_actual}")
    except Exception as e:
        return f"Error IA: {str(e)}"

@app.route("/api/ia", methods=["POST"])
def api_ia():
    if not GROQ_API_KEY: return jsonify({"respuesta": "Error: API Key no configurada."})
    
    pregunta = request.json.get("pregunta", "")
    try:
        client = Groq(api_key=GROQ_API_KEY)
        completion = client.chat.completions.create(
            model="openai/gpt-oss-20b", 
            messages=[{"role": "system", "content": "Eres un contador chileno. Responde corto."}, {"role": "user", "content": pregunta}],
            temperature=0.5, max_tokens=500
        )
        return jsonify({"respuesta": completion.choices[0].message.content})
    except Exception as e:
        return jsonify({"respuesta": f"Error: {str(e)}"})

@app.route("/api/get_rut")
@login_required
def get_rut():
    uid = session['user_id']
    nombre = request.args.get("nombre", "").strip()
    con = get_db(); cur = con.cursor()
    cur.execute("SELECT rut FROM docs WHERE nombre LIKE %s AND user_id=%s ORDER BY id DESC LIMIT 1", (f"%{nombre}%", uid))
    row = cur.fetchone(); cur.close(); con.close()
    return jsonify({"rut": row[0] if row and row[0] else ""})

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port)